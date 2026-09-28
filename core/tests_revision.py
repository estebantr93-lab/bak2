"""Casos reportados en la revisión del compañero: cada test reproduce la falla y verifica la corrección."""
import datetime
import io

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from actividades.models import Activity, SocialCase
from core.models import ActivityType, Period
from core.testing import SesionTestMixin, sembrar_datos_demo
from evidencias.models import Evidence
from funcionarios.models import Employee


class RevisionCompaneroTests(SesionTestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        sembrar_datos_demo()
        cls.ana = Employee.objects.get(name='Ana Pérez (Centro)')
        cls.cerrado = Period.objects.get(is_closed=True)
        cls.abierto = Period.objects.get(is_closed=False)
        tipo = ActivityType.objects.get(code='ATC-01')
        # Actividad de Ana en el período cerrado (histórica, creada antes del cierre).
        cls.historica = Activity.objects.create(
            number='HIST-001', employee=cls.ana, delegation=cls.ana.delegation, period=cls.cerrado,
            activity_type=tipo, date=datetime.date(2026, 2, 10), description='Histórica', evidence_code='EV-HIST-001',
        )
        # Actividad de otro funcionario de Centro (María, admin de Centro).
        cls.maria = Employee.objects.get(name='María Soto (Admin Centro)')
        cls.ajena = Activity.objects.create(
            number='AJENA-001', employee=cls.maria, delegation=cls.maria.delegation, period=cls.abierto,
            activity_type=tipo, date=datetime.date(2026, 7, 1), description='De María', evidence_code='EV-AJENA-001',
        )

    def _datos(self, actividad, **extra):
        datos = {
            'number': actividad.number, 'employee': actividad.employee_id, 'period': actividad.period_id or '',
            'activity_type': actividad.activity_type_id, 'date': actividad.date.isoformat(),
            'description': actividad.description, 'evidence_code': actividad.evidence_code,
        }
        datos.update(extra)
        return datos

    # 1) Período cerrado
    def test_no_se_puede_sacar_una_actividad_del_periodo_cerrado_dejandolo_vacio(self):
        self.ingresar('funcionario_centro')
        url = reverse('actividad_update', args=[self.historica.pk])
        response = self.client.post(url, self._datos(self.historica, period=''))
        self.assertEqual(response.status_code, 403)
        self.historica.refresh_from_db()
        self.assertEqual(self.historica.period, self.cerrado)

    def test_actividad_de_periodo_cerrado_no_se_edita_ni_elimina_ni_como_admin(self):
        self.ingresar('admin_centro')
        self.assertEqual(self.client.get(reverse('actividad_update', args=[self.historica.pk])).status_code, 403)
        self.assertEqual(self.client.post(reverse('actividad_delete', args=[self.historica.pk])).status_code, 403)
        self.assertIsNone(Activity.all_objects.get(pk=self.historica.pk).deleted_at)

    def test_periodo_es_obligatorio_en_el_formulario_web(self):
        self.ingresar('funcionario_centro')
        propia = Activity.objects.filter(employee=self.ana, period=self.abierto).exclude(validation_status='approved').first()
        response = self.client.post(reverse('actividad_update', args=[propia.pk]), self._datos(propia, period=''))
        self.assertIn('period', response.context['form'].errors)

    def test_el_modelo_impide_cambiar_de_periodo_una_actividad_cerrada(self):
        from django.core.exceptions import ValidationError

        self.historica.period = self.abierto
        with self.assertRaises(ValidationError):
            self.historica.full_clean()

    # 2) Opción (a): ve toda su delegación, modifica solo lo suyo
    def test_funcionario_ve_actividades_ajenas_de_su_delegacion(self):
        self.ingresar('funcionario_centro')
        response = self.client.get(reverse('actividad_list'), {'page_size': 30})
        self.assertIn(self.ajena, list(Activity.objects.filter(delegation=self.ana.delegation)))
        fila = next((f for f in response.context['rows'] if f['obj'].pk == self.ajena.pk), None)
        if fila:  # si está en la primera página, se ve pero sin botones de edición
            self.assertFalse(fila['editable'])

    def test_funcionario_no_puede_apropiarse_de_una_actividad_ajena(self):
        self.ingresar('funcionario_centro')
        url = reverse('actividad_update', args=[self.ajena.pk])
        self.assertEqual(self.client.get(url).status_code, 403)
        response = self.client.post(url, self._datos(self.ajena, employee=self.ana.pk, description='apropiada'))
        self.assertEqual(response.status_code, 403)
        self.ajena.refresh_from_db()
        self.assertEqual(self.ajena.employee, self.maria)

    def test_funcionario_si_edita_sus_actividades(self):
        self.ingresar('funcionario_centro')
        propia = Activity.objects.filter(employee=self.ana, period=self.abierto).exclude(validation_status='approved').first()
        response = self.client.post(reverse('actividad_update', args=[propia.pk]), self._datos(propia, description='ok'))
        self.assertRedirects(response, reverse('actividad_list'))

    def test_admin_de_delegacion_si_edita_actividades_de_su_equipo(self):
        self.ingresar('admin_centro')
        response = self.client.get(reverse('actividad_update', args=[Activity.objects.filter(
            employee=self.ana, period=self.abierto).first().pk]))
        self.assertEqual(response.status_code, 200)

    def test_funcionario_solo_adjunta_evidencias_a_sus_actividades(self):
        self.ingresar('funcionario_centro')
        response = self.client.get(reverse('evidencia_create'))
        opciones = set(response.context['form'].fields['activity'].queryset.values_list('employee_id', flat=True))
        self.assertEqual(opciones, {self.ana.pk})

    # 3) Excel: sin fórmulas inyectadas
    def test_excel_guarda_como_texto_lo_que_empieza_con_igual(self):
        Activity.objects.filter(pk=self.ajena.pk).update(number='=HYPERLINK("http://x","clic")')
        self.ingresar('admin_centro')
        hoja = load_workbook(io.BytesIO(self.client.get(reverse('actividad_export')).content)).active
        celdas = [c for fila in hoja.iter_rows(min_row=2) for c in fila if str(c.value).startswith('=HYPERLINK')]
        self.assertTrue(celdas)
        self.assertTrue(all(c.data_type == 's' for c in celdas))

    # 4) Gestión social eliminada: se puede volver a registrar el mismo número
    def test_gestion_eliminada_se_reactiva_al_registrar_el_mismo_numero(self):
        self.ingresar('funcionario_centro')
        gestion = SocialCase.objects.filter(activity__employee=self.ana).first()
        actividad, paso = gestion.activity, gestion.step_number
        gestion.delete()
        response = self.client.post(reverse('atencion_create'), {
            'activity': actividad.pk, 'step_number': paso, 'description': 'Nueva gestión',
        })
        self.assertRedirects(response, reverse('atencion_list'))
        activa = SocialCase.objects.get(activity=actividad, step_number=paso)
        self.assertEqual(activa.description, 'Nueva gestión')
        self.assertEqual(SocialCase.all_objects.filter(activity=actividad, step_number=paso).count(), 1)

    # 5) restore() recupera los hijos eliminados junto con el padre
    def test_restore_recupera_los_hijos_eliminados_con_el_padre(self):
        actividad = Activity.objects.filter(employee=self.ana, period=self.abierto).first()
        previa = Evidence.objects.create(activity=actividad, description='eliminada antes')
        previa.delete()
        junto = Evidence.objects.create(activity=actividad, description='cae con el padre')
        actividad.delete()
        self.assertIsNotNone(Evidence.all_objects.get(pk=junto.pk).deleted_at)
        Activity.all_objects.get(pk=actividad.pk).restore()
        self.assertIsNone(Evidence.all_objects.get(pk=junto.pk).deleted_at)
        self.assertIsNotNone(Evidence.all_objects.get(pk=previa.pk).deleted_at)  # esta ya estaba eliminada


class PoliticaUnicaEnElAdminTests(SesionTestMixin, TestCase):
    """La misma regla de modificación rige en el CRUD web y en el Django Admin (todas las vías)."""

    @classmethod
    def setUpTestData(cls):
        sembrar_datos_demo()
        cls.ana = Employee.objects.get(name='Ana Pérez (Centro)')
        cls.maria = Employee.objects.get(name='María Soto (Admin Centro)')
        cls.abierto = Period.objects.get(is_closed=False)
        cls.cerrado = Period.objects.get(is_closed=True)
        tipo = ActivityType.objects.get(code='ATC-01')
        cls.ajena = Activity.objects.create(
            number='ADM-AJENA', employee=cls.maria, delegation=cls.maria.delegation, period=cls.abierto,
            activity_type=tipo, date=datetime.date(2026, 7, 1), description='De María', evidence_code='EV-ADM-AJ',
        )
        cls.propia = Activity.objects.filter(employee=cls.ana, period=cls.abierto).exclude(validation_status='approved').first()
        cls.cerrada = Activity.objects.create(
            number='ADM-CERR', employee=cls.ana, delegation=cls.ana.delegation, period=cls.cerrado,
            activity_type=tipo, date=datetime.date(2026, 2, 1), description='Cerrada', evidence_code='EV-ADM-CE',
        )

    def _post_admin(self, actividad, **extra):
        datos = {
            'number': actividad.number, 'employee': actividad.employee_id, 'delegation': actividad.delegation_id,
            'period': actividad.period_id, 'activity_type': actividad.activity_type_id,
            'date': actividad.date.isoformat(), 'description': 'cambiada', 'evidence_code': actividad.evidence_code,
            'validation_status': actividad.validation_status,
            'evidence_items-TOTAL_FORMS': 0, 'evidence_items-INITIAL_FORMS': 0,
        }
        datos.update(extra)
        return self.client.post(f'/admin/actividades/activity/{actividad.pk}/change/', datos)

    def test_admin_funcionario_ve_la_actividad_ajena_en_solo_lectura_y_no_puede_guardarla(self):
        self.ingresar('funcionario_centro')
        response = self.client.get(f'/admin/actividades/activity/{self.ajena.pk}/change/')
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['has_change_permission'])
        self.assertEqual(self._post_admin(self.ajena, employee=self.ana.pk).status_code, 403)
        self.ajena.refresh_from_db()
        self.assertEqual((self.ajena.employee, self.ajena.description), (self.maria, 'De María'))

    def test_admin_funcionario_si_modifica_la_suya(self):
        self.ingresar('funcionario_centro')
        response = self.client.get(f'/admin/actividades/activity/{self.propia.pk}/change/')
        self.assertTrue(response.context['has_change_permission'])

    def test_admin_periodo_cerrado_no_se_edita_ni_elimina_ni_siquiera_el_superusuario(self):
        for username in ('admin_centro', 'admin_sgr'):
            self.ingresar(username)
            response = self.client.get(f'/admin/actividades/activity/{self.cerrada.pk}/change/')
            self.assertFalse(response.context['has_change_permission'], username)
            self.assertEqual(self.client.post(
                f'/admin/actividades/activity/{self.cerrada.pk}/delete/', {'post': 'yes'}).status_code, 403, username)
            self.assertIsNone(Activity.all_objects.get(pk=self.cerrada.pk).deleted_at)

    def test_admin_borrado_masivo_omite_lo_que_no_se_puede_modificar(self):
        self.ingresar('admin_centro')
        self.client.post('/admin/actividades/activity/', {
            'action': 'delete_selected', '_selected_action': [self.cerrada.pk], 'post': 'yes',
        })
        self.assertIsNone(Activity.all_objects.get(pk=self.cerrada.pk).deleted_at)

    def test_admin_opciones_y_autocompletado_solo_ofrecen_lo_propio(self):
        self.ingresar('funcionario_centro')
        # Formulario de alta de evidencia en el Admin: el campo 'activity' solo acepta actividades propias.
        response = self.client.get('/admin/evidencias/evidence/add/')
        opciones = response.context['adminform'].form.fields['activity'].queryset
        self.assertEqual(set(opciones.values_list('employee_id', flat=True)), {self.ana.pk})
        # El autocompletado de ese campo tampoco ofrece actividades ajenas.
        response = self.client.get('/admin/autocomplete/', {
            'term': 'ADM', 'app_label': 'evidencias', 'model_name': 'evidence', 'field_name': 'activity',
        })
        numeros = {item['text'] for item in response.json()['results']}
        self.assertNotIn('ADM-AJENA', numeros)


class AlcanceDeLlavesForaneasEnElAdminTests(SesionTestMixin, TestCase):
    """Ningún campo relacionado de ningún Admin acotado ofrece datos de otra delegación."""

    @classmethod
    def setUpTestData(cls):
        sembrar_datos_demo()

    def test_ningun_campo_relacionado_ofrece_datos_de_otra_delegacion(self):
        from django.contrib import admin
        from django.contrib.auth.models import User
        from django.db import models
        from django.test import RequestFactory

        from core.admin_utils import ScopedModelAdmin

        for username, ajena in [('admin_centro', 'Delegación Norte'), ('admin_norte', 'Delegación Centro'),
                                ('funcionario_centro', 'Delegación Norte')]:
            request = RequestFactory().get('/')
            request.user = User.objects.get(username=username)
            for modelo, model_admin in admin.site._registry.items():
                if not isinstance(model_admin, ScopedModelAdmin):
                    continue
                for campo in modelo._meta.fields:
                    if not isinstance(campo, models.ForeignKey):
                        continue
                    opciones = model_admin.formfield_for_foreignkey(campo, request).queryset
                    relacionado = admin.site._registry.get(campo.related_model)
                    if isinstance(relacionado, ScopedModelAdmin):
                        ruta = 'name' if relacionado.scope_by == 'pk' else f'{relacionado.scope_by}__name'
                        fuera = opciones.filter(**{ruta: ajena})
                    elif campo.related_model is User:
                        fuera = opciones.exclude(employee__delegation=request.user.employee.delegation)
                    else:
                        continue
                    self.assertFalse(fuera.exists(), f'{username}: {modelo.__name__}.{campo.name}')

    def test_admin_centro_no_valida_evidencias_de_norte_y_el_revisor_es_el_propio_usuario(self):
        from evidencias.models import Validation

        self.ingresar('admin_centro')
        norte = Evidence.objects.filter(activity__delegation__name='Delegación Norte').first()
        centro = Evidence.objects.filter(activity__delegation__name='Delegación Centro').first()
        response = self.client.post('/admin/evidencias/validation/add/', {
            'evidence': norte.pk, 'status': 'approved', 'comment': 'intento',
        })
        self.assertEqual(response.status_code, 200)  # formulario con error: la opción no es válida
        self.assertFalse(Validation.objects.filter(evidence=norte, comment='intento').exists())
        response = self.client.post('/admin/evidencias/validation/add/', {
            'evidence': centro.pk, 'status': 'approved', 'comment': 'ok centro',
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Validation.objects.get(comment='ok centro').reviewer.username, 'admin_centro')
