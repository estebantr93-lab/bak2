import datetime
import io

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from actividades.models import Activity, SocialCase
from agenda.models import Commitment
from core.testing import SesionTestMixin, sembrar_datos_demo
from evidencias.models import Evidence, Validation
from funcionarios.models import Employee

CRUDS = {
    # prefijo: (modelo, camino a la delegación)
    'actividad': (Activity, 'delegation'),
    'atencion': (SocialCase, 'activity__delegation'),
    'evidencia': (Evidence, 'activity__delegation'),
    'compromiso': (Commitment, 'delegation'),
}


class BaseCrud(SesionTestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        sembrar_datos_demo()


class SeguridadComunTests(BaseCrud):
    def test_anonimo_va_al_login_en_todos_los_crud(self):
        for prefijo in CRUDS:
            for accion in ('list', 'create', 'export'):
                url = reverse(f'{prefijo}_{accion}')
                response = self.client.get(url)
                self.assertRedirects(response, f"{reverse('login')}?next={url}", msg_prefix=url)

    def test_listados_respetan_scoping_de_delegacion(self):
        self.ingresar('admin_centro')
        for prefijo, (modelo, camino) in CRUDS.items():
            response = self.client.get(reverse(f'{prefijo}_list'), {'page_size': 30})
            self.assertEqual(response.status_code, 200, prefijo)
            delegaciones = {
                modelo.objects.filter(pk=f['obj'].pk).values_list(f'{camino}__name', flat=True).get()
                for f in response.context['rows']
            }
            self.assertTrue(delegaciones <= {'Delegación Centro'}, (prefijo, delegaciones))

    def test_objeto_de_otra_delegacion_da_404_al_editar_y_eliminar(self):
        self.ingresar('admin_centro')
        for prefijo, (modelo, camino) in CRUDS.items():
            ajeno = modelo.objects.filter(**{f'{camino}__name': 'Delegación Norte'}).first()
            self.assertIsNotNone(ajeno, prefijo)
            self.assertEqual(self.client.get(reverse(f'{prefijo}_update', args=[ajeno.pk])).status_code, 404)
            self.assertEqual(self.client.post(reverse(f'{prefijo}_delete', args=[ajeno.pk])).status_code, 404)
            self.assertIsNone(modelo.all_objects.get(pk=ajeno.pk).deleted_at)

    def test_eliminar_es_por_post_y_logico(self):
        self.ingresar('admin_centro')
        for prefijo, (modelo, camino) in CRUDS.items():
            obj = modelo.objects.filter(**{f'{camino}__name': 'Delegación Centro'}).first()
            url = reverse(f'{prefijo}_delete', args=[obj.pk])
            self.assertEqual(self.client.get(url).status_code, 405)
            self.assertRedirects(self.client.post(url), reverse(f'{prefijo}_list'))
            self.assertIsNotNone(modelo.all_objects.get(pk=obj.pk).deleted_at, prefijo)
            self.assertNotIn(obj, list(self.client.get(reverse(f'{prefijo}_list')).context['page_obj']))

    def test_funcionario_no_puede_eliminar_en_ningun_crud(self):
        self.ingresar('funcionario_centro')
        for prefijo, (modelo, camino) in CRUDS.items():
            obj = modelo.objects.filter(**{f'{camino}__name': 'Delegación Centro'}).first()
            self.assertEqual(self.client.post(reverse(f'{prefijo}_delete', args=[obj.pk])).status_code, 403)

    def test_verificador_sin_permiso_sobre_compromisos_recibe_403(self):
        self.ingresar('verificador_leia')
        self.assertEqual(self.client.get(reverse('compromiso_list')).status_code, 403)
        self.assertEqual(self.client.get(reverse('compromiso_export')).status_code, 403)


class ExportacionExcelTests(BaseCrud):
    def _libro(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment;', response['Content-Disposition'])
        return load_workbook(io.BytesIO(response.content)).active

    def test_excel_tiene_encabezados_y_solo_datos_del_alcance(self):
        self.ingresar('admin_centro')
        hoja = self._libro(self.client.get(reverse('actividad_export')))
        encabezados = [c.value for c in hoja[1]]
        self.assertEqual(encabezados[:3], ['Número', 'Fecha', 'Funcionario'])
        esperadas = Activity.objects.filter(delegation__name='Delegación Centro')
        self.assertEqual(hoja.max_row - 1, esperadas.count())
        numeros = {fila[0] for fila in hoja.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(numeros, set(esperadas.values_list('number', flat=True)))

    def test_excel_excluye_eliminados(self):
        self.ingresar('admin_centro')
        eliminada = Activity.objects.filter(delegation__name='Delegación Centro').first()
        eliminada.delete()
        hoja = self._libro(self.client.get(reverse('actividad_export')))
        numeros = {fila[0] for fila in hoja.iter_rows(min_row=2, values_only=True)}
        self.assertNotIn(eliminada.number, numeros)

    def test_todos_los_crud_exportan(self):
        self.ingresar('admin_sgr')
        for prefijo, (modelo, _) in CRUDS.items():
            hoja = self._libro(self.client.get(reverse(f'{prefijo}_export')))
            self.assertEqual(hoja.max_row - 1, modelo.objects.count(), prefijo)


class CompromisoTests(BaseCrud):
    def datos(self, **extra):
        responsable = Employee.objects.get(name='Ana Pérez (Centro)')
        datos = {
            'title': 'Reunión con junta de vecinos', 'description': '', 'responsible': responsable.pk,
            'due_date': (timezone.localdate() + datetime.timedelta(days=10)).isoformat(),
            'status': 'pending', 'notes': '',
        }
        datos.update(extra)
        return datos

    def test_admin_centro_crea_compromiso_en_su_delegacion(self):
        self.ingresar('admin_centro')
        response = self.client.post(reverse('compromiso_create'), self.datos(), follow=True)
        self.assertContains(response, 'registrado correctamente')
        creado = Commitment.objects.get(title='Reunión con junta de vecinos')
        self.assertEqual(creado.delegation.name, 'Delegación Centro')

    def test_fecha_en_el_pasado_se_rechaza(self):
        self.ingresar('admin_centro')
        response = self.client.post(reverse('compromiso_create'), self.datos(due_date='2020-01-01'))
        self.assertIn('due_date', response.context['form'].errors)
        self.assertTrue(response.context['modal_abierto'])

    def test_realizado_exige_observaciones(self):
        self.ingresar('admin_centro')
        response = self.client.post(reverse('compromiso_create'), self.datos(status='done'))
        self.assertIn('notes', response.context['form'].errors)

    def test_responsable_debe_ser_de_la_misma_delegacion(self):
        self.ingresar('admin_sgr')
        norte = Employee.objects.get(name='Carlos Rojas (Norte)')
        centro = Employee.objects.get(name='Ana Pérez (Centro)').delegation
        response = self.client.post(
            reverse('compromiso_create'), self.datos(responsible=norte.pk, delegation=centro.pk),
        )
        self.assertIn('responsible', response.context['form'].errors)

    def test_titulo_demasiado_corto(self):
        self.ingresar('admin_centro')
        response = self.client.post(reverse('compromiso_create'), self.datos(title='  ab '))
        self.assertIn('title', response.context['form'].errors)


class AtencionSocialTests(BaseCrud):
    def setUp(self):
        self.social = Activity.objects.filter(
            delegation__name='Delegación Centro', activity_type__category='social',
        ).first()

    def test_rango_de_gestion_1_a_3(self):
        self.ingresar('admin_centro')
        response = self.client.post(reverse('atencion_create'), {
            'activity': self.social.pk, 'step_number': 4, 'description': 'Cuarta gestión',
        })
        self.assertIn('step_number', response.context['form'].errors)

    def test_gestion_duplicada_se_rechaza(self):
        self.ingresar('admin_centro')
        existente = SocialCase.objects.filter(activity=self.social).first()
        response = self.client.post(reverse('atencion_create'), {
            'activity': self.social.pk, 'step_number': existente.step_number, 'description': 'Repetida',
        })
        self.assertFalse(response.context['form'].is_valid())

    def test_solo_actividades_de_tipo_social(self):
        self.ingresar('admin_sgr')
        no_social = Activity.objects.exclude(activity_type__category='social').first()
        response = self.client.post(reverse('atencion_create'), {
            'activity': no_social.pk, 'step_number': 2, 'description': 'x',
        })
        self.assertIn('activity', response.context['form'].errors)

    def test_crear_segunda_gestion(self):
        self.ingresar('funcionario_centro')
        libre = next(n for n in (1, 2, 3) if not SocialCase.objects.filter(activity=self.social, step_number=n).exists())
        response = self.client.post(reverse('atencion_create'), {
            'activity': self.social.pk, 'step_number': libre, 'description': 'Seguimiento en terreno', 'date': timezone.localdate().isoformat(),
        })
        self.assertRedirects(response, reverse('atencion_list'))


class EvidenciaRevisionTests(BaseCrud):
    def test_verificador_aprueba_y_queda_registro(self):
        self.ingresar('verificador_leia')
        evidencia = Evidence.objects.filter(status='pending').first()
        response = self.client.get(reverse('evidencia_update', args=[evidencia.pk]))
        self.assertIn('status', response.context['form'].fields)
        datos = {'activity': evidencia.activity_id, 'description': evidencia.description,
                 'status': 'approved', 'result': 'Documento conforme'}
        # Las evidencias del seed no traen archivo; se asigna uno para que la edición no exija adjuntarlo.
        if not evidencia.file:
            evidencia.file.name = 'evidencias/prueba.pdf'
            evidencia.save(update_fields=['file'])
        response = self.client.post(reverse('evidencia_update', args=[evidencia.pk]), datos)
        self.assertRedirects(response, reverse('evidencia_list'))
        evidencia.refresh_from_db()
        self.assertEqual(evidencia.status, 'approved')
        self.assertEqual(evidencia.reviewed_by.username, 'verificador_leia')
        self.assertTrue(Validation.objects.filter(evidence=evidencia, status='approved').exists())

    def test_funcionario_no_ve_el_campo_estado(self):
        self.ingresar('funcionario_centro')
        response = self.client.get(reverse('evidencia_list'))
        self.assertNotIn('status', response.context['form'].fields)
        self.assertEqual(User.objects.get(username='funcionario_centro').has_perm('evidencias.change_evidence'), False)


class ListadoInteractivoTests(BaseCrud):
    """Pestañas con conteo, búsqueda, orden por columna y exportación de las filas seleccionadas."""

    def test_pestanas_y_tarjetas_cuentan_solo_el_alcance_del_usuario(self):
        self.ingresar('admin_centro')
        response = self.client.get(reverse('actividad_list'))
        pestanas = {p['clave']: p for p in response.context['pestanas']}
        centro = Activity.objects.filter(delegation__name='Delegación Centro')
        self.assertEqual(pestanas['todas']['total'], centro.count())
        self.assertTrue(pestanas['todas']['activa'])
        for clave, estado in (('pendientes', 'pending'), ('aprobadas', 'approved'), ('rechazadas', 'rejected')):
            self.assertEqual(pestanas[clave]['total'], centro.filter(validation_status=estado).count(), clave)
        self.assertEqual([t['clave'] for t in response.context['tarjetas']],
                         ['todas', 'pendientes', 'aprobadas', 'rechazadas'])

    def test_pestana_activa_filtra_el_listado(self):
        self.ingresar('admin_centro')
        response = self.client.get(reverse('actividad_list'), {'status': 'rejected', 'page_size': 100})
        activa = [p['clave'] for p in response.context['pestanas'] if p['activa']]
        self.assertEqual(activa, ['rechazadas'])
        self.assertTrue(all(f['obj'].validation_status == 'rejected' for f in response.context['rows']))

    def test_todos_los_listados_tienen_pestanas_y_responden_para_cada_rol(self):
        for usuario in ('admin_sgr', 'admin_centro', 'verificador_leia', 'funcionario_centro'):
            self.ingresar(usuario)
            for prefijo in CRUDS:
                response = self.client.get(reverse(f'{prefijo}_list'))
                if response.status_code == 403:
                    continue  # el rol no tiene ese módulo (lo cubre test_verificador_sin_permiso...)
                self.assertEqual(response.status_code, 200, (usuario, prefijo))
                self.assertTrue(response.context['pestanas'], (usuario, prefijo))

    def test_busqueda_por_numero_y_nombre_dentro_del_alcance(self):
        self.ingresar('admin_centro')
        actividad = Activity.objects.filter(delegation__name='Delegación Centro').first()
        response = self.client.get(reverse('actividad_list'), {'q': actividad.number})
        self.assertEqual([f['obj'].pk for f in response.context['rows']], [actividad.pk])
        ajena = Activity.objects.filter(delegation__name='Delegación Norte').first()
        response = self.client.get(reverse('actividad_list'), {'q': ajena.number})
        self.assertEqual(response.context['rows'], [])

    def test_busqueda_se_mantiene_en_pestanas_y_paginacion(self):
        self.ingresar('admin_centro')
        response = self.client.get(reverse('actividad_list'), {'q': 'Ana'})
        self.assertIn('q=Ana', response.context['filters_query'])
        self.assertTrue(all('q=Ana' in p['url'] for p in response.context['pestanas']))

    def test_orden_por_columna_permitida(self):
        self.ingresar('admin_sgr')
        response = self.client.get(reverse('compromiso_list'), {'orden': '-due_date', 'page_size': 100})
        fechas = [f['obj'].due_date for f in response.context['rows']]
        self.assertEqual(fechas, sorted(fechas, reverse=True))

    def test_orden_fuera_de_la_lista_blanca_se_ignora(self):
        self.ingresar('admin_sgr')
        for orden in ('delegation__employees__user__password', 'pk; DROP', '-description'):
            response = self.client.get(reverse('actividad_list'), {'orden': orden})
            self.assertEqual(response.status_code, 200, orden)
            self.assertEqual(response.context['orden'], '', orden)

    def test_exportar_seleccion_solo_incluye_ids_del_alcance(self):
        self.ingresar('admin_centro')
        propias = list(Activity.objects.filter(delegation__name='Delegación Centro')[:2])
        ajena = Activity.objects.filter(delegation__name='Delegación Norte').first()
        ids = ','.join(str(a.pk) for a in [*propias, ajena])
        response = self.client.get(reverse('actividad_export'), {'ids': ids})
        self.assertEqual(response.status_code, 200)
        hoja = load_workbook(io.BytesIO(response.content)).active
        numeros = {fila[0] for fila in hoja.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(numeros, {a.number for a in propias})

    def test_ids_invalidos_no_rompen_la_exportacion(self):
        self.ingresar('admin_centro')
        response = self.client.get(reverse('actividad_export'), {'ids': 'abc,-3,'})
        self.assertEqual(response.status_code, 200)
        hoja = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(hoja.max_row, 1)  # solo encabezados: una selección inválida no exporta todo


class AdminEtiquetasDeEstadoTests(BaseCrud):
    def test_changelist_muestra_el_estado_como_etiqueta_y_ordena_por_el_campo(self):
        self.ingresar('admin_sgr')
        response = self.client.get(reverse('admin:actividades_activity_changelist'), {'o': '6'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<span class="estado estado-pending">Pendiente</span>', html=True)
        self.assertContains(response, 'Estado de validación')


class PanelDeFiltrosTests(BaseCrud):
    """Filtros propios de cada módulo: validados por lista blanca y siempre dentro del alcance."""

    def _pks(self, response):
        return {f['obj'].pk for f in response.context['rows']}

    def _todas(self, nombre, params):
        return self.client.get(reverse(nombre), {**params, 'page_size': 30})

    def test_evidencias_por_ultima_modificacion(self):
        self.ingresar('admin_sgr')
        antigua, reciente = Evidence.objects.all()[:2]
        Evidence.all_objects.exclude(pk=reciente.pk).update(updated_at=timezone.now() - datetime.timedelta(days=45))
        response = self._todas('evidencia_list', {'modificada': '7'})
        self.assertEqual(self._pks(response), {reciente.pk})
        self.assertIn('última modificación: últimos 7 días', response.context['filters_desc'])
        response = self._todas('evidencia_list', {'modificada': 'mas_30'})
        self.assertIn(antigua.pk, self._pks(response))
        self.assertNotIn(reciente.pk, self._pks(response))

    def test_ultima_modificacion_existe_en_todos_los_modulos(self):
        self.ingresar('admin_sgr')
        for prefijo, param in (('actividad', 'modificada'), ('atencion', 'modificada'),
                               ('evidencia', 'modificada'), ('compromiso', 'modificado')):
            campos = self.client.get(reverse(f'{prefijo}_list')).context['panel']['campos']
            self.assertIn(param, [c['filtro'].param for c in campos], prefijo)

    def test_valores_no_permitidos_se_ignoran(self):
        self.ingresar('admin_sgr')
        total = self.client.get(reverse('evidencia_list')).context['page_obj'].paginator.count
        for params in ({'modificada': '99'}, {'delegacion': '999999'}, {'registrada_desde': '2026-13-45'},
                       {'archivo': 'quizas'}, {'tipo': 'x'}):
            response = self.client.get(reverse('evidencia_list'), params)
            self.assertEqual(response.status_code, 200, params)
            self.assertEqual(response.context['page_obj'].paginator.count, total, params)
            self.assertEqual(response.context['chips'], [], params)

    def test_opciones_del_panel_respetan_el_alcance(self):
        self.ingresar('admin_centro')
        response = self.client.get(reverse('actividad_list'))
        campos = {c['filtro'].param: c['filtro'] for c in response.context['panel']['campos']}
        self.assertNotIn('delegacion', campos)  # una sola delegación: el filtro no aporta
        nombres = {texto for _, texto in campos['funcionario'].opciones}
        self.assertEqual(nombres, set(Employee.objects.filter(delegation__name='Delegación Centro').values_list('name', flat=True)))
        # Pedir otra delegación por URL no abre sus datos.
        norte = Activity.objects.filter(delegation__name='Delegación Norte').first().delegation_id
        response = self.client.get(reverse('actividad_list'), {'delegacion': norte, 'page_size': 30})
        self.assertTrue(all(f['obj'].delegation.name == 'Delegación Centro' for f in response.context['rows']))

    def test_superadmin_filtra_por_delegacion_y_tipo(self):
        self.ingresar('admin_sgr')
        actividad = Activity.objects.filter(delegation__name='Delegación Norte').first()
        response = self._todas('actividad_list', {'delegacion': actividad.delegation_id, 'tipo': actividad.activity_type_id})
        esperadas = Activity.objects.filter(delegation=actividad.delegation, activity_type=actividad.activity_type)
        self.assertEqual(response.context['page_obj'].paginator.count, esperadas.count())

    def test_rango_de_vencimiento_en_compromisos(self):
        self.ingresar('admin_sgr')
        desde = timezone.localdate()
        hasta = desde + datetime.timedelta(days=30)
        response = self._todas('compromiso_list', {'vence_desde': desde.isoformat(), 'vence_hasta': hasta.isoformat()})
        esperados = Commitment.objects.filter(due_date__range=(desde, hasta))
        self.assertEqual(response.context['page_obj'].paginator.count, esperados.count())
        self.assertTrue(response.context['chips'][0]['texto'].startswith('vencimiento desde'))

    def test_revisores_solo_del_alcance(self):
        self.ingresar('admin_centro')
        campos = {c['filtro'].param: c['filtro'] for c in self.client.get(reverse('evidencia_list')).context['panel']['campos']}
        revisores = set(Evidence.objects.filter(activity__delegation__name='Delegación Centro', reviewed_by__isnull=False)
                        .values_list('reviewed_by__username', flat=True))
        self.assertEqual({texto for _, texto in campos['revisor'].opciones}, revisores)

    def test_filtro_de_archivo(self):
        self.ingresar('admin_sgr')
        con = self._todas('evidencia_list', {'archivo': 'con'}).context['page_obj'].paginator.count
        sin = self._todas('evidencia_list', {'archivo': 'sin'}).context['page_obj'].paginator.count
        self.assertEqual(con + sin, Evidence.objects.count())
        self.assertEqual(sin, Evidence.objects.filter(file='').count() + Evidence.objects.filter(file__isnull=True).count())

    def test_chip_quita_solo_su_filtro_y_pestanas_lo_conservan(self):
        self.ingresar('admin_sgr')
        response = self.client.get(reverse('evidencia_list'), {'modificada': '7', 'status': 'pending', 'q': 'EVI'})
        chips = {c['texto']: c['url'] for c in response.context['chips']}
        url = chips['última modificación: últimos 7 días']
        self.assertNotIn('modificada=', url)
        self.assertIn('status=pending', url)
        self.assertIn('q=EVI', url)
        self.assertTrue(all('modificada=7' in p['url'] for p in response.context['pestanas']))

    def test_excel_aplica_los_filtros_del_panel(self):
        self.ingresar('admin_sgr')
        reciente = Evidence.objects.first()
        Evidence.all_objects.exclude(pk=reciente.pk).update(updated_at=timezone.now() - datetime.timedelta(days=45))
        response = self.client.get(reverse('evidencia_export'), {'modificada': 'hoy'})
        hoja = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual([fila[0] for fila in hoja.iter_rows(min_row=2, values_only=True)], [reciente.unique_code])

    def test_enlace_del_dashboard_por_periodo_sigue_funcionando(self):
        self.ingresar('admin_sgr')
        actividad = Activity.objects.first()
        response = self._todas('evidencia_list', {'period': actividad.period_id})
        esperadas = Evidence.objects.filter(activity__period_id=actividad.period_id)
        self.assertEqual(response.context['page_obj'].paginator.count, esperadas.count())

    def test_admin_filtra_por_ultima_modificacion(self):
        self.ingresar('admin_sgr')
        response = self.client.get(reverse('admin:evidencias_evidence_changelist'))
        self.assertContains(response, 'última modificación')
        filtro = next(f for f in response.context['cl'].filter_specs if f.title == 'última modificación')
        hoy = next(c for c in filtro.choices(response.context['cl']) if c['display'] == 'Hoy')
        response = self.client.get(reverse('admin:evidencias_evidence_changelist') + hoy['query_string'])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['cl'].result_count, Evidence.objects.count())  # los datos de prueba son de hoy
