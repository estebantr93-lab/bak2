"""Catálogo de reglas de negocio del SGR, verificadas contra el código (una prueba por regla).

Cada prueba ejerce la regla por la vía real (formulario web, Admin o modelo) con los datos de demo.
Correr solo este catálogo:  python manage.py test core.tests_reglas_negocio -v 2
"""
import datetime
import io
import shutil
import tempfile
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image

from actividades.models import Activity, SocialCase
from agenda.forms import CommitmentForm
from agenda.models import Commitment, CommitmentFollowUp
from colaboracion.models import AuditLog
from core.models import ActivityType, Period, Position
from core.testing import CLAVE_TEST, sembrar_datos_demo
from evidencias.models import Evidence, Validation
from evidencias.services import registrar_revision
from funcionarios.models import Employee, PasswordResetCode
from medicion.forms import GoalForm
from medicion.models import Goal
from medicion.services import calcular_cumplimiento_ponderado, calcular_semaforo

MEDIA_TEMPORAL = tempfile.mkdtemp()


def png():
    buffer = io.BytesIO()
    Image.new('RGB', (20, 20), (173, 0, 0)).save(buffer, format='PNG')
    return buffer.getvalue()


@override_settings(MEDIA_ROOT=MEDIA_TEMPORAL,
                   MAILERS={'default': {'BACKEND': 'django.core.mail.backends.locmem.EmailBackend'}})
class ReglasDeNegocioTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        sembrar_datos_demo()
        cls.abierto = Period.objects.get(is_closed=False)
        cls.cerrado = Period.objects.get(is_closed=True)
        cls.ana = Employee.objects.get(user__username='funcionario_centro')
        cls.carlos = Employee.objects.get(user__username='funcionario_norte')
        cls.tipo = ActivityType.objects.get(code='ATC-01')
        cls.social = ActivityType.objects.get(code='SOC-04')

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA_TEMPORAL, ignore_errors=True)

    def ingresar(self, username):
        self.client.logout()
        self.assertTrue(self.client.login(username=username, password=CLAVE_TEST))

    def actividad(self, empleado=None, **extra):
        empleado = empleado or self.ana
        n = Activity.all_objects.count() + 1
        datos = dict(number=f'RN-{n}', employee=empleado, delegation=empleado.delegation, period=self.abierto,
                     activity_type=self.tipo, date=self.abierto.start_date, description='Regla',
                     evidence_code=f'EV-RN-{n}')
        datos.update(extra)
        return Activity.objects.create(**datos)

    def datos_web(self, **extra):
        n = Activity.all_objects.count() + 1
        datos = {'number': f'WEB-{n}', 'employee': self.ana.pk, 'period': self.abierto.pk,
                 'activity_type': self.tipo.pk, 'date': self.abierto.start_date.isoformat(),
                 'description': 'Desde la web', 'evidence_code': f'EV-WEB-{n}'}
        datos.update(extra)
        return datos

    # ---------------- Acceso, roles y alcance ----------------
    def test_rn_login_rechaza_cuentas_sin_rol(self):
        User.objects.create_user('sin_rol', password=CLAVE_TEST)
        response = self.client.post(reverse('login'), {'username': 'sin_rol', 'password': CLAVE_TEST})
        self.assertContains(response, 'no tiene un rol asignado')

    def test_rn_login_rechaza_funcionario_desactivado(self):
        Employee.objects.filter(pk=self.ana.pk).update(is_active=False)
        response = self.client.post(reverse('login'), {'username': 'funcionario_centro', 'password': CLAVE_TEST})
        self.assertContains(response, 'desactivado')

    def test_rn_superadmin_y_verificador_ven_ambas_delegaciones(self):
        for username in ('admin_sgr', 'verificador_leia'):
            self.ingresar(username)
            nombres = {s['delegation'].name for s in self.client.get(reverse('dashboard')).context['secciones']}
            self.assertEqual(nombres, {'Delegación Centro', 'Delegación Norte'}, username)

    def test_rn_admin_de_delegacion_solo_ve_y_modifica_su_delegacion(self):
        norte = self.actividad(self.carlos)
        self.ingresar('admin_centro')
        self.assertEqual(self.client.get(reverse('actividad_update', args=[norte.pk])).status_code, 404)
        filas = self.client.get(reverse('actividad_list'), {'page_size': 30}).context['rows']
        self.assertEqual({f['obj'].delegation.name for f in filas}, {'Delegación Centro'})

    def test_rn_funcionario_ve_su_delegacion_pero_solo_modifica_lo_propio(self):
        maria = Employee.objects.get(user__username='admin_centro')
        ajena = self.actividad(maria)
        self.ingresar('funcionario_centro')
        self.assertEqual(self.client.get(reverse('actividad_update', args=[ajena.pk])).status_code, 403)
        propia = self.actividad()
        self.assertEqual(self.client.get(reverse('actividad_update', args=[propia.pk])).status_code, 200)

    def test_rn_admin_de_delegacion_no_gestiona_usuarios_ni_mueve_funcionarios(self):
        self.ingresar('admin_centro')
        self.assertIn(self.client.get('/admin/auth/user/').status_code, (302, 403))
        ficha = self.client.get(f'/admin/funcionarios/employee/{self.ana.pk}/change/')
        self.assertNotIn('delegation', ficha.context['adminform'].form.fields)
        self.assertNotIn('user', ficha.context['adminform'].form.fields)

    def test_rn_funcionario_no_se_elimina_se_desactiva(self):
        self.ingresar('admin_sgr')
        self.assertEqual(self.client.get(f'/admin/funcionarios/employee/{self.ana.pk}/delete/').status_code, 403)

    # ---------------- Contraseñas y recuperación ----------------
    def test_rn_politica_de_contrasenas(self):
        from django.contrib.auth.password_validation import validate_password

        for debil in ('Corta#1a', 'sinmayuscula#2026', 'SINMINUSCULA#2026', 'SinNumero#Clave', 'SinEspecial2026x'):
            with self.assertRaises(ValidationError, msg=debil):
                validate_password(debil)
        validate_password('Correcta#2026x')

    def test_rn_recuperacion_codigo_6_digitos_hash_unico_uso_y_tope(self):
        from funcionarios.recuperacion import CODIGO_OK, solicitar_codigo, validar_codigo
        from django.core import mail

        correo = self.ana.user.email
        registro = solicitar_codigo(correo)
        codigo = mail.outbox[-1].body.split(': ')[1][:6]
        self.assertTrue(codigo.isdigit() and len(codigo) == 6)
        self.assertNotEqual(registro.code_hash, codigo)
        self.assertEqual(validar_codigo(correo, codigo)[0], CODIGO_OK)
        self.assertNotEqual(validar_codigo(correo, codigo)[0], CODIGO_OK)  # un solo uso
        for _ in range(10):
            solicitar_codigo(correo)
        self.assertLessEqual(PasswordResetCode.objects.filter(user=self.ana.user).count(), 5)

    # ---------------- Actividades ----------------
    def test_rn_periodo_obligatorio_y_fecha_dentro_del_periodo_y_no_futura(self):
        self.ingresar('funcionario_centro')
        for extra, campo in [({'period': ''}, 'period'),
                             ({'date': (timezone.localdate() + datetime.timedelta(days=1)).isoformat()}, 'date'),
                             ({'date': (self.abierto.start_date - datetime.timedelta(days=1)).isoformat()}, 'date')]:
            response = self.client.post(reverse('actividad_create'), self.datos_web(**extra))
            self.assertIn(campo, response.context['form'].errors, extra)

    def test_rn_delegacion_de_la_actividad_es_la_del_funcionario(self):
        actividad = Activity(number='X', employee=self.ana, delegation=self.carlos.delegation, period=self.abierto,
                             activity_type=self.tipo, date=self.abierto.start_date, description='x', evidence_code='EV-X')
        with self.assertRaises(ValidationError):
            actividad.full_clean()

    def test_rn_numero_de_actividad_unico_incluso_contra_eliminadas(self):
        eliminada = self.actividad()
        eliminada.delete()
        self.ingresar('funcionario_centro')
        response = self.client.post(reverse('actividad_create'), self.datos_web(number=eliminada.number.lower()))
        self.assertIn('number', response.context['form'].errors)

    def test_rn_estado_de_la_actividad_se_deriva_de_sus_evidencias(self):
        verificador = User.objects.get(username='verificador_leia')
        actividad = self.actividad()
        e1 = Evidence.objects.create(activity=actividad, description='1', file='evidencias/1.png')
        registrar_revision(e1, verificador, 'rejected', 'Borrosa')
        actividad.refresh_from_db()
        self.assertEqual(actividad.validation_status, 'rejected')
        e2 = Evidence.objects.create(activity=actividad, description='2', file='evidencias/2.png')
        actividad.refresh_from_db()
        self.assertEqual(actividad.validation_status, 'pending')
        registrar_revision(e2, verificador, 'approved', 'Correcta')
        actividad.refresh_from_db()
        self.assertEqual(actividad.validation_status, 'approved')

    def test_rn_actividad_aprobada_congelada_para_el_funcionario(self):
        aprobada = self.actividad(activity_type=self.social, validation_status='approved')
        gestion = SocialCase.objects.create(activity=aprobada, step_number=1, description='g')
        self.ingresar('funcionario_centro')
        self.assertEqual(self.client.get(reverse('actividad_update', args=[aprobada.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse('atencion_update', args=[gestion.pk])).status_code, 403)
        form = self.client.get(reverse('evidencia_list')).context['form']
        self.assertNotIn(aprobada, form.fields['activity'].queryset)
        self.ingresar('admin_centro')
        self.assertEqual(self.client.get(reverse('actividad_update', args=[aprobada.pk])).status_code, 200)

    # ---------------- Período cerrado ----------------
    def test_rn_periodo_cerrado_congela_todo_incluso_para_el_superadmin(self):
        historica = self.actividad(period=self.cerrado, date=self.cerrado.start_date)
        evidencia = Evidence.objects.create(activity=historica, description='h', file='evidencias/h.png')
        for username in ('admin_centro', 'admin_sgr'):
            self.ingresar(username)
            self.assertEqual(self.client.get(reverse('actividad_update', args=[historica.pk])).status_code, 403)
            self.assertEqual(self.client.post(reverse('evidencia_delete', args=[evidencia.pk])).status_code, 403)
        response = self.client.post(reverse('actividad_create'), self.datos_web(period=self.cerrado.pk))
        self.assertIn('period', response.context['form'].errors)
        with self.assertRaises(ValidationError):
            Evidence(activity=historica, description='nueva').full_clean()

    def test_rn_metas_de_un_periodo_cerrado_no_se_modifican(self):
        meta = Goal.objects.filter(period=self.abierto).first()
        meta.period = self.cerrado
        with self.assertRaises(ValidationError):
            meta.full_clean()

    # ---------------- Evidencias ----------------
    def test_rn_evidencia_exige_archivo_valido_jpg_png_pdf_de_hasta_2mb(self):
        self.ingresar('funcionario_centro')
        actividad = self.actividad()
        for nombre, contenido in [('a.exe', b'MZ'), ('a.png', b'no es imagen'), ('a.pdf', b'no es pdf'),
                                  ('a.pdf', b'%PDF-' + b'0' * (2 * 1024 * 1024 + 1))]:
            response = self.client.post(reverse('evidencia_create'), {
                'activity': actividad.pk, 'description': nombre, 'file': SimpleUploadedFile(nombre, contenido)})
            self.assertIn('file', response.context['form'].errors, nombre)
        response = self.client.post(reverse('evidencia_create'), {'activity': actividad.pk, 'description': 'sin archivo'})
        self.assertIn('file', response.context['form'].errors)
        response = self.client.post(reverse('evidencia_create'), {
            'activity': actividad.pk, 'description': 'ok', 'file': SimpleUploadedFile('mi foto.png', png())})
        self.assertEqual(response.status_code, 302)
        self.assertNotIn('mi foto', Evidence.objects.get(description='ok').file.name)  # nombre seguro

    def test_rn_revisar_requiere_permiso_y_rechazar_exige_motivo(self):
        evidencia = Evidence.objects.create(activity=self.actividad(), description='r', file='evidencias/r.png')
        self.ingresar('funcionario_centro')
        self.assertEqual(self.client.get(reverse('evidencia_update', args=[evidencia.pk])).status_code, 403)
        self.ingresar('verificador_leia')
        response = self.client.post(reverse('evidencia_update', args=[evidencia.pk]), {'status': 'rejected', 'result': ''})
        self.assertIn('result', response.context['form'].errors)

    def test_rn_revision_deja_revisor_validacion_y_traza(self):
        evidencia = Evidence.objects.create(activity=self.actividad(), description='t', file='evidencias/t.png')
        self.ingresar('verificador_leia')
        self.client.post(reverse('evidencia_update', args=[evidencia.pk]), {'status': 'approved', 'result': 'Bien'})
        evidencia.refresh_from_db()
        self.assertEqual(evidencia.reviewed_by.username, 'verificador_leia')
        self.assertTrue(Validation.objects.filter(evidence=evidencia, status='approved').exists())
        self.assertTrue(AuditLog.objects.filter(entity_type='Evidence', entity_id=evidencia.pk).exists())

    def test_rn_verificador_solo_cambia_estado_y_resultado(self):
        actividades = [self.actividad(), self.actividad()]
        evidencia = Evidence.objects.create(activity=actividades[0], description='v', file='evidencias/v.png')
        self.ingresar('verificador_leia')
        self.client.post(reverse('evidencia_update', args=[evidencia.pk]),
                         {'activity': actividades[1].pk, 'description': 'cambiada', 'status': 'approved', 'result': 'ok'})
        evidencia.refresh_from_db()
        self.assertEqual((evidencia.activity_id, evidencia.description), (actividades[0].pk, 'v'))

    # ---------------- Atención social ----------------
    def test_rn_atencion_social_solo_en_actividades_sociales_y_maximo_3_gestiones(self):
        no_social = self.actividad()
        with self.assertRaises(ValidationError):
            SocialCase(activity=no_social, step_number=1, description='x').full_clean()
        social = self.actividad(activity_type=self.social)
        for paso in (1, 2, 3):
            SocialCase.objects.create(activity=social, step_number=paso, description=f'g{paso}')
        with self.assertRaises(ValidationError):
            SocialCase(activity=social, step_number=4, description='x').full_clean()

    # ---------------- Compromisos ----------------
    def test_rn_compromiso_titulo_vencimiento_responsable_y_realizado(self):
        admin = User.objects.get(username='admin_sgr')
        datos = {'title': 'Abc', 'delegation': self.ana.delegation_id, 'responsible': self.carlos.pk,
                 'due_date': (timezone.localdate() - datetime.timedelta(days=1)).isoformat(), 'status': 'done', 'notes': ''}
        form = CommitmentForm(data=datos, user=admin)
        self.assertFalse(form.is_valid())
        self.assertIn('title', form.errors)       # mínimo 5 caracteres
        self.assertIn('due_date', form.errors)    # no nace vencido
        self.assertIn('responsible', form.errors)  # de la misma delegación
        # Con responsable válido, marcarlo realizado sin observaciones también se rechaza.
        datos.update(title='Título válido', responsible=self.ana.pk, due_date=timezone.localdate().isoformat())
        form = CommitmentForm(data=datos, user=admin)
        self.assertFalse(form.is_valid())
        self.assertIn('notes', form.errors)

    def test_rn_seguimiento_actualiza_el_estado_del_compromiso(self):
        compromiso = Commitment.objects.exclude(status='done').first()
        CommitmentFollowUp.objects.create(commitment=compromiso, new_status='in_progress', description='Avanza')
        compromiso.refresh_from_db()
        self.assertEqual(compromiso.status, 'in_progress')

    # ---------------- Metas y cumplimiento ----------------
    def test_rn_metas_ponderadores_hasta_100_y_valores_positivos_y_unicas(self):
        cargo = Position.objects.create(name='Cargo de prueba')
        Goal.objects.create(position=cargo, period=self.abierto, activity_type=self.tipo, target=5, weight=Decimal('70'))
        form = GoalForm(data={'position': cargo.pk, 'period': self.abierto.pk, 'activity_type': self.social.pk,
                              'target': 0, 'weight': '40'})
        self.assertFalse(form.is_valid())
        self.assertIn('target', form.errors)
        form = GoalForm(data={'position': cargo.pk, 'period': self.abierto.pk, 'activity_type': self.social.pk,
                              'target': 3, 'weight': '40'})
        self.assertFalse(form.is_valid())  # 70 + 40 > 100
        form = GoalForm(data={'position': cargo.pk, 'period': self.abierto.pk, 'activity_type': self.tipo.pk,
                              'target': 3, 'weight': '10'})
        self.assertFalse(form.is_valid())  # misma meta repetida

    def test_rn_005_cumplimiento_por_tipo_ponderado_con_tope_del_periodo(self):
        from types import SimpleNamespace

        metas = [SimpleNamespace(activity_type_id=1, activity_type='A', target=10, weight=Decimal('50')),
                 SimpleNamespace(activity_type_id=2, activity_type='B', target=10, weight=Decimal('50'))]
        pct, _ = calcular_cumplimiento_ponderado({1: 40, 3: 99}, metas, tope=self.abierto.max_cap)
        self.assertEqual(pct, Decimal('75'))  # 400 % topado a 150 %; el tipo 3 no tiene meta
        self.assertEqual(calcular_semaforo(Decimal('80'), Decimal('70')), 'green')
        self.assertEqual(calcular_semaforo(Decimal('50'), Decimal('70')), 'amber')
        self.assertEqual(calcular_semaforo(Decimal('30'), Decimal('70')), 'red')

    def test_rn_periodos_sin_solapamiento_e_inicio_antes_del_termino(self):
        with self.assertRaises(ValidationError):
            Period(name='Solapado', start_date=self.abierto.start_date, end_date=self.abierto.end_date).full_clean()
        with self.assertRaises(ValidationError):
            Period(name='Invertido', start_date=datetime.date(2030, 5, 1), end_date=datetime.date(2030, 1, 1)).full_clean()

    # ---------------- Borrado lógico y auditoría ----------------
    def test_rn_borrado_logico_en_cascada_y_lo_ve_el_superadmin(self):
        actividad = self.actividad()
        evidencia = Evidence.objects.create(activity=actividad, description='c', file='evidencias/c.png')
        self.ingresar('admin_centro')
        self.assertEqual(self.client.get(reverse('actividad_delete', args=[actividad.pk])).status_code, 405)
        self.client.post(reverse('actividad_delete', args=[actividad.pk]))
        self.assertIsNotNone(Evidence.all_objects.get(pk=evidencia.pk).deleted_at)
        self.assertTrue(AuditLog.objects.filter(action='eliminar', entity_id=actividad.pk, user__username='admin_centro').exists())
        self.ingresar('admin_sgr')
        self.assertEqual(self.client.get(f'/admin/actividades/activity/{actividad.pk}/change/').status_code, 200)

    def test_rn_traza_de_auditoria_de_solo_lectura(self):
        self.ingresar('admin_sgr')
        traza = AuditLog.objects.create(action='x', entity_type='x')
        self.assertEqual(self.client.get('/admin/colaboracion/auditlog/add/').status_code, 403)
        self.assertEqual(self.client.post(f'/admin/colaboracion/auditlog/{traza.pk}/delete/', {'post': 'yes'}).status_code, 403)

    # ---------------- Reportes ----------------
    def test_rn_excel_respeta_alcance_y_neutraliza_formulas(self):
        Commitment.objects.create(title='=HYPERLINK("x")', delegation=self.ana.delegation,
                                  due_date=timezone.localdate(), status='registered')
        self.ingresar('admin_centro')
        hoja = load_workbook(io.BytesIO(self.client.get(reverse('compromiso_export')).content)).active
        celdas = [fila[0] for fila in hoja.iter_rows(min_row=2)]
        formula = next(c for c in celdas if str(c.value).startswith('='))
        self.assertEqual(formula.data_type, 's')  # texto, no fórmula
        self.assertEqual({fila[1] for fila in hoja.iter_rows(min_row=2, values_only=True)}, {'Delegación Centro'})
