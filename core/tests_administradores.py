"""Lo que hace cada administrador (general, Centro y Norte) de punta a punta, y la regla de negocio:
lo que eliminan los administradores de delegación lo sigue viendo el administrador general."""
import io
import os
import shutil
import tempfile
import zipfile

from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image

from actividades.models import Activity
from colaboracion.models import AuditLog
from core.models import Period
from core.testing import CLAVE_TEST, sembrar_datos_demo
from django.core.files.uploadedfile import SimpleUploadedFile
from evidencias.archivos import pdf_de_texto
from evidencias.models import Evidence
from funcionarios.models import Employee

MEDIA_TEMPORAL = tempfile.mkdtemp()
ADMINISTRADORES = {'admin_sgr': 'Delegación Centro', 'admin_centro': 'Delegación Centro', 'admin_norte': 'Delegación Norte'}


def foto(formato):
    buffer = io.BytesIO()
    Image.new('RGB', (64, 48), (173, 0, 0)).save(buffer, format=formato)
    return buffer.getvalue()


@override_settings(MEDIA_ROOT=MEDIA_TEMPORAL)
class AdministradoresDePuntaAPuntaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        sembrar_datos_demo()
        cls.periodo = Period.objects.get(is_closed=False)

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA_TEMPORAL, ignore_errors=True)

    def ingresar(self, username):
        response = self.client.post(reverse('login'), {'username': username, 'password': CLAVE_TEST})
        self.assertRedirects(response, reverse('dashboard'), msg_prefix=username)

    def registrar_actividad(self, username, delegacion):
        funcionario = Employee.objects.filter(delegation__name=delegacion, is_active=True).first()
        hoy = timezone.localdate()
        numero = f'ADM-{username}'.upper()
        response = self.client.post(reverse('actividad_create'), {
            'number': numero, 'employee': funcionario.pk, 'period': self.periodo.pk,
            'activity_type': Activity.objects.first().activity_type_id, 'date': hoy.isoformat(),
            'description': f'Registrada por {username}', 'evidence_code': f'EV-{numero}',
        })
        self.assertEqual(response.status_code, 302, username)
        actividad = Activity.objects.get(number=numero)
        self.assertEqual((actividad.date, actividad.delegation.name), (hoy, delegacion))
        self.assertTrue(self.periodo.start_date <= actividad.date <= self.periodo.end_date)
        return actividad

    def subir(self, actividad, nombre, contenido):
        response = self.client.post(reverse('evidencia_create'), {
            'activity': actividad.pk, 'description': nombre, 'file': SimpleUploadedFile(nombre, contenido),
        })
        self.assertEqual(response.status_code, 302, nombre)
        return Evidence.objects.get(activity=actividad, description=nombre)

    def test_cada_administrador_registra_sube_archivos_y_descarga_excel(self):
        for username, delegacion in ADMINISTRADORES.items():
            with self.subTest(username=username):
                self.ingresar(username)
                self.assertEqual(self.client.get(reverse('dashboard')).status_code, 200)
                actividad = self.registrar_actividad(username, delegacion)

                for nombre, contenido in [('foto.jpg', foto('JPEG')), ('foto.png', foto('PNG')),
                                          ('acta.pdf', pdf_de_texto(['Acta de prueba']))]:
                    evidencia = self.subir(actividad, nombre, contenido)
                    self.assertTrue(os.path.exists(evidencia.file.path))
                    self.assertEqual(timezone.localtime(evidencia.registered_at).date(), timezone.localdate())
                    if nombre != 'acta.pdf':
                        Image.open(evidencia.file.path).verify()  # la foto subida se abre

                lista = self.client.get(reverse('evidencia_list'), {'activity': actividad.pk}).content.decode()
                self.assertIn(f'{timezone.localdate():%d-%m-%Y}', lista)  # fecha en formato chileno

                for url in ('actividad_export', 'atencion_export', 'evidencia_export', 'compromiso_export'):
                    respuesta = self.client.get(reverse(url))
                    self.assertEqual(respuesta.status_code, 200, url)
                    zipfile.ZipFile(io.BytesIO(respuesta.content)).testzip()
                    hoja = load_workbook(io.BytesIO(respuesta.content)).active
                    if url == 'actividad_export':
                        numeros = [fila[0] for fila in hoja.iter_rows(min_row=2, values_only=True)]
                        self.assertIn(actividad.number, numeros)
                        delegaciones = {fila[3] for fila in hoja.iter_rows(min_row=2, values_only=True)}
                        if username != 'admin_sgr':
                            self.assertEqual(delegaciones, {delegacion})  # cada admin, solo su delegación
                self.client.logout()

    def test_lo_que_elimina_un_admin_de_delegacion_lo_sigue_viendo_el_superadmin(self):
        for username, delegacion in [('admin_centro', 'Delegación Centro'), ('admin_norte', 'Delegación Norte')]:
            with self.subTest(username=username):
                self.ingresar(username)
                actividad = self.registrar_actividad(username, delegacion)
                evidencia = self.subir(actividad, 'foto.png', foto('PNG'))
                otra = self.subir(actividad, 'otra.png', foto('PNG'))
                self.assertEqual(self.client.post(reverse('evidencia_delete', args=[evidencia.pk])).status_code, 302)
                self.assertEqual(self.client.post(reverse('actividad_delete', args=[actividad.pk])).status_code, 302)
                # Para el admin de delegación ya no existe (ni en la web ni en el Admin).
                self.assertEqual(self.client.get(f'/admin/actividades/activity/{actividad.pk}/change/').status_code, 302)
                self.client.logout()

                self.ingresar('admin_sgr')
                eliminadas = self.client.get('/admin/actividades/activity/', {'registro': 'eliminados'})
                self.assertIn(actividad, list(eliminadas.context['cl'].result_list))
                ficha = self.client.get(f'/admin/actividades/activity/{actividad.pk}/change/')
                self.assertEqual(ficha.status_code, 200)
                self.assertFalse(ficha.context['has_change_permission'])  # se consulta, no se edita
                evidencias = self.client.get('/admin/evidencias/evidence/', {'registro': 'eliminados'})
                self.assertTrue({evidencia, otra} <= set(evidencias.context['cl'].result_list))
                self.assertTrue(os.path.exists(Evidence.all_objects.get(pk=evidencia.pk).file.path))  # el archivo se conserva
                # Queda quién lo eliminó.
                self.assertTrue(AuditLog.objects.filter(action='eliminar', entity_type='Activity', entity_id=actividad.pk,
                                                        user__username=username).exists())
                self.assertTrue(AuditLog.objects.filter(action='eliminar', entity_type='Evidence', entity_id=evidencia.pk,
                                                        user__username=username).exists())
                # Y puede restaurarla: vuelve con las evidencias que cayeron junto con ella.
                self.client.post('/admin/actividades/activity/', {
                    'action': 'restaurar_registros', '_selected_action': [actividad.pk], 'registro': 'eliminados',
                })
                self.assertTrue(Activity.objects.filter(pk=actividad.pk).exists())
                self.assertTrue(Evidence.objects.filter(pk=otra.pk).exists())
                self.assertFalse(Evidence.objects.filter(pk=evidencia.pk).exists())  # esa se había eliminado antes
                self.client.logout()

    def test_no_se_restaura_un_registro_cuyo_padre_sigue_eliminado(self):
        self.ingresar('admin_centro')
        actividad = self.registrar_actividad('admin_centro', 'Delegación Centro')
        evidencia = self.subir(actividad, 'foto.png', foto('PNG'))
        self.client.post(reverse('evidencia_delete', args=[evidencia.pk]))
        self.client.post(reverse('actividad_delete', args=[actividad.pk]))
        self.client.logout()

        self.ingresar('admin_sgr')
        restaurar_evidencia = {'action': 'restaurar_registros', '_selected_action': [evidencia.pk], 'registro': 'eliminados'}
        respuesta = self.client.post('/admin/evidencias/evidence/', restaurar_evidencia, follow=True)
        # La actividad sigue eliminada: la evidencia no vuelve (quedaría colgando de algo que nadie ve).
        self.assertFalse(Evidence.objects.filter(pk=evidencia.pk).exists())
        self.assertIn('Restaure primero', ' '.join(str(m) for m in respuesta.context['messages']))
        self.assertFalse(AuditLog.objects.filter(action='restaurar', entity_type='Evidence', entity_id=evidencia.pk).exists())

        # Restaurada la actividad, la evidencia (eliminada antes y por separado) ya se puede restaurar.
        self.client.post('/admin/actividades/activity/', {
            'action': 'restaurar_registros', '_selected_action': [actividad.pk], 'registro': 'eliminados',
        })
        self.assertFalse(Evidence.objects.filter(pk=evidencia.pk).exists())
        self.client.post('/admin/evidencias/evidence/', restaurar_evidencia)
        self.assertTrue(Evidence.objects.filter(pk=evidencia.pk).exists())
        self.client.logout()

        # La lista web del admin de delegación y el dashboard vuelven a coincidir.
        self.ingresar('admin_centro')
        self.assertEqual(self.client.get(reverse('evidencia_list'), {'activity': actividad.pk}).context['page_obj'].paginator.count, 1)

    def test_gestiones_y_seguimientos_tampoco_vuelven_sin_su_padre(self):
        from django.core.exceptions import ValidationError

        from actividades.models import SocialCase
        from agenda.models import CommitmentFollowUp

        for hijo in (SocialCase.objects.first(), CommitmentFollowUp.objects.first()):
            with self.subTest(modelo=type(hijo).__name__):
                padre = hijo.activity if isinstance(hijo, SocialCase) else hijo.commitment
                hijo.delete()
                padre.delete()
                hijo = type(hijo).all_objects.get(pk=hijo.pk)
                with self.assertRaisesMessage(ValidationError, 'Restaure primero'):
                    hijo.restore()
                self.assertIsNotNone(type(hijo).all_objects.get(pk=hijo.pk).deleted_at)
                padre.restore()
                hijo.restore()
                self.assertTrue(type(hijo).objects.filter(pk=hijo.pk).exists())

    def test_los_admins_de_delegacion_no_ven_eliminados_ni_pueden_restaurar(self):
        self.ingresar('admin_centro')
        response = self.client.get('/admin/actividades/activity/')
        self.assertNotIn('restaurar_registros', response.context['action_form'].fields['action'].choices.__repr__())
        self.assertNotIn('estado del registro', [str(f.title) for f in response.context['cl'].filter_specs])
        eliminada = Activity.objects.filter(delegation__name='Delegación Centro').first()
        eliminada.delete()
        self.assertNotIn(eliminada, list(response.context['cl'].queryset))
        self.assertEqual(self.client.get(f'/admin/actividades/activity/{eliminada.pk}/change/').status_code, 302)
