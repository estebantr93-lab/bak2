"""Exportación de reportes a Excel (.xlsx) con openpyxl.

openpyxl arma el libro en memoria: Workbook() crea el archivo, hoja.append(fila) agrega filas y
libro.save(respuesta) lo escribe directamente en la respuesta HTTP (no queda archivo en disco).
Quien llama entrega las filas ya filtradas, así el reporte respeta permisos, scoping y borrado lógico.
"""
from datetime import date, datetime

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def valor_excel(valor, como_texto):
    """Fechas y números se guardan como tipos nativos de Excel; el resto como texto."""
    if isinstance(valor, datetime) and timezone.is_aware(valor):
        valor = timezone.localtime(valor).replace(tzinfo=None)  # Excel no guarda zona horaria
    if isinstance(valor, (int, float, date)) and not isinstance(valor, bool):
        return valor
    return como_texto(valor)


def respuesta_xlsx(titulo, encabezados, filas, nombre_archivo):
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]  # Excel limita el nombre de la hoja a 31 caracteres
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='AD0000')
    for fila in filas:
        hoja.append(fila)
    for i, encabezado in enumerate(encabezados, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = max(12, len(encabezado) + 4)
    hoja.freeze_panes = 'A2'
    respuesta = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    libro.save(respuesta)
    return respuesta
