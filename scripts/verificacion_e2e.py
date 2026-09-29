"""Verificación de punta a punta de los puntos pedidos (rúbrica y funciones del sistema).

Recorre el sistema como cada rol sobre la base de desarrollo (con seed_data --volumen cargado):
login y recuperación, sesiones y paginación, aislamiento por delegación, los 4 CRUD con archivos,
Excel, borrado lógico, revisión de evidencias, período cerrado, dashboard por rol, Admin,
volumen de datos, nombres de tablas y secretos fuera del repositorio.

No deja cambios: todo corre dentro de una transacción que se revierte y con MEDIA_ROOT temporal.
Uso, desde la raíz del proyecto:  python manage.py shell < scripts/verificacion_e2e.py
"""
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile

from django.conf import settings
from django.contrib.auth.models import User
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection, transaction
from django.test import Client, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image

from actividades.models import Activity, SocialCase
from agenda.models import Commitment, CommitmentFollowUp
from colaboracion.models import AuditLog
from core.models import Period
from evidencias.models import Evidence, Validation
from funcionarios.models import Employee, PasswordResetCode

from django.test.utils import setup_test_environment

setup_test_environment()  # habilita response.context y el correo en memoria
RESULTADOS = []


def ok(punto, condicion, detalle=''):
    RESULTADOS.append((punto, bool(condicion), detalle))


def cliente(username):
    c = Client()
    c.force_login(User.objects.get(username=username))
    return c


def png():
    b = io.BytesIO()
    Image.new('RGB', (40, 30), (173, 0, 0)).save(b, format='PNG')
    return b.getvalue()


def xlsx(resp):
    zipfile.ZipFile(io.BytesIO(resp.content)).testzip()
    return load_workbook(io.BytesIO(resp.content)).active


media = tempfile.mkdtemp()
CLAVE = 'Verif#2026Clave'
correo_locmem = {'default': {'BACKEND': 'django.core.mail.backends.locmem.EmailBackend'}}

with override_settings(MEDIA_ROOT=media, MAILERS=correo_locmem, ALLOWED_HOSTS=['testserver']), transaction.atomic():
    for u in User.objects.filter(username__in=['admin_sgr', 'admin_centro', 'admin_norte', 'funcionario_centro',
                                                'funcionario_norte', 'verificador_leia']):
        u.set_password(CLAVE)
        u.save()
    # Los intentos de ingreso anteriores (por ejemplo, de una demo reciente) no deben bloquear las cuentas
    # de esta verificación. Todo lo que se hace aquí se revierte al final (set_rollback).
    AuditLog.objects.filter(entity_type='Acceso').delete()
    abierto = Period.objects.get(is_closed=False)
    cerrado = Period.objects.get(is_closed=True)
    centro_act = Activity.objects.filter(delegation__name='Delegación Centro', period=abierto).first()
    norte_act = Activity.objects.filter(delegation__name='Delegación Norte', period=abierto).first()

    # ---------------- 1. Login, logout y roles ----------------
    c = Client()
    ok('Login: página disponible', c.get(reverse('login')).status_code == 200)
    r = c.post(reverse('login'), {'username': 'admin_centro', 'password': CLAVE})
    ok('Login: credenciales válidas llevan al dashboard', r.status_code == 302 and r.url == reverse('dashboard'), r.get('Location'))
    r = Client().post(reverse('login'), {'username': 'admin_centro', 'password': 'mala'})
    ok('Login: error genérico con clave incorrecta', 'Usuario o contraseña incorrectos' in r.content.decode())
    sin_rol = User.objects.create_user('verif_sin_rol', password=CLAVE)
    r = Client().post(reverse('login'), {'username': 'verif_sin_rol', 'password': CLAVE})
    ok('Login: cuenta sin rol rechazada', 'no tiene un rol asignado' in r.content.decode())
    ok('Logout: GET no cierra sesión (solo POST)', c.get(reverse('logout')).status_code == 405)
    r = c.post(reverse('logout'))
    ok('Logout: POST cierra sesión y vuelve al login', r.status_code == 302 and reverse('login') in r.url)
    ok('Anónimo: dashboard redirige al login', Client().get(reverse('dashboard')).status_code == 302)
    ok('/admin/login/ redirige al login del sistema', reverse('login') in Client().get('/admin/login/').get('Location', ''))
    # Límite de intentos (OWASP A07) y traza de accesos (A09), con una cuenta que no se usa después.
    for _ in range(settings.LOGIN_MAX_INTENTOS):
        Client().post(reverse('login'), {'username': 'admin_norte', 'password': 'mala'}, REMOTE_ADDR='10.20.30.40')
    r = Client().post(reverse('login'), {'username': 'admin_norte', 'password': CLAVE}, REMOTE_ADDR='10.20.30.40')
    ok('Login: tras 5 fallos se bloquea aunque la clave sea correcta', 'Demasiados intentos fallidos' in r.content.decode())
    ok('Traza: ingresos fallidos y bloqueo con IP', AuditLog.objects.filter(
        action='login_fallido', detail='admin_norte', ip='10.20.30.40').count() == settings.LOGIN_MAX_INTENTOS
        and AuditLog.objects.filter(action='login_bloqueado', detail='admin_norte').exists())
    AuditLog.objects.filter(entity_type='Acceso', detail='admin_norte').delete()  # la cuenta se usa más adelante
    ok('CSP: la respuesta declara Content-Security-Policy', "default-src 'self'" in
       Client().get(reverse('login')).headers.get('Content-Security-Policy', ''))

    # ---------------- 2. Recuperación por código de 6 dígitos ----------------
    mail.outbox = []
    c = Client()
    correo = User.objects.get(username='funcionario_centro').email
    c.post(reverse('recuperar_solicitar'), {'email': correo})
    codigo = re.search(r'\b(\d{6})\b', mail.outbox[-1].body).group(1) if mail.outbox else None
    ok('Recuperación: envía un código de 6 dígitos por correo', codigo is not None)
    ok('Recuperación: guarda solo el hash del código', codigo and not PasswordResetCode.objects.filter(code_hash=codigo).exists())
    r = Client().post(reverse('recuperar_solicitar'), {'email': 'nadie@noexiste.cl'})
    ok('Recuperación: respuesta genérica si el correo no existe', r.status_code == 302)
    r = c.post(reverse('recuperar_codigo'), {'code': '000000' if codigo != '000000' else '111111'})
    ok('Recuperación: código incorrecto se rechaza', 'incorrecto' in r.content.decode())
    r = c.post(reverse('recuperar_codigo'), {'code': codigo})
    ok('Recuperación: código correcto avanza', r.status_code == 302 and r.url == reverse('recuperar_nueva'))
    r = c.post(reverse('recuperar_nueva'), {'new_password1': 'debil', 'new_password2': 'debil'})
    ok('Política de contraseñas: rechaza una clave débil', r.status_code == 200)
    nueva = 'NuevaClave#2026x'
    r = c.post(reverse('recuperar_nueva'), {'new_password1': nueva, 'new_password2': nueva})
    ok('Recuperación: define la nueva contraseña', r.status_code == 302 and Client().login(username='funcionario_centro', password=nueva))
    r = Client().post(reverse('recuperar_codigo'), {'code': codigo})
    ok('Recuperación: el código es de un solo uso', r.status_code == 302 and 'recuperar' in r.url)

    # ---------------- 3. Sesiones ----------------
    ok('Sesión: expira a las 2 horas y la cookie es HttpOnly', settings.SESSION_COOKIE_AGE == 7200 and settings.SESSION_COOKIE_HTTPONLY)
    c = cliente('admin_centro')
    c.get(reverse('actividad_list'), {'page_size': 5})
    r = c.get(reverse('evidencia_list'))
    ok('Sesión: la paginación elegida (5) se recuerda en todos los listados', r.context['page_obj'].paginator.per_page == 5)
    c.get(reverse('actividad_list'), {'page_size': 999})
    ok('Paginación: valor no permitido se ignora', c.session.get('page_size') == 5)
    for n in (15, 30):
        r = c.get(reverse('actividad_list'), {'page_size': n})
        ok(f'Paginación: {n} por página', len(r.context['page_obj'].object_list) == n)
    c.get(reverse('dashboard'), {'period': cerrado.pk})
    ok('Sesión: el período del dashboard se recuerda', c.get(reverse('dashboard')).context['period'] == cerrado)

    # ---------------- 4. Aislamiento por delegación ----------------
    c = cliente('admin_centro')
    ok('Aislamiento: admin_centro no abre una actividad de Norte (404)',
       c.get(reverse('actividad_update', args=[norte_act.pk])).status_code == 404)
    r = c.get(reverse('actividad_list'), {'page_size': 30})
    ok('Aislamiento: el listado de admin_centro es solo de Centro',
       {f['obj'].delegation.name for f in r.context['rows']} == {'Delegación Centro'})
    hoja = xlsx(c.get(reverse('actividad_export')))
    ok('Aislamiento: el Excel de admin_centro es solo de Centro',
       {fila[3] for fila in hoja.iter_rows(min_row=2, values_only=True)} == {'Delegación Centro'})
    ok('Aislamiento: admin_centro no ve Norte en el Admin',
       c.get(f'/admin/actividades/activity/{norte_act.pk}/change/').status_code in (302, 404))
    cn = cliente('admin_norte')
    ok('Aislamiento: admin_norte no abre una actividad de Centro (404)',
       cn.get(reverse('actividad_update', args=[centro_act.pk])).status_code == 404)
    cf = cliente('funcionario_centro')
    ajena = Activity.objects.filter(delegation__name='Delegación Centro', period=abierto).exclude(
        employee__user__username='funcionario_centro').first()
    ok('Propiedad: el funcionario no edita actividades ajenas (403)',
       cf.get(reverse('actividad_update', args=[ajena.pk])).status_code == 403)
    ok('Permisos: el funcionario no ve el Admin de usuarios', cf.get('/admin/auth/user/').status_code in (302, 403))
    cv = cliente('verificador_leia')
    ok('Permisos: el verificador no abre Compromisos (403)', cv.get(reverse('compromiso_list')).status_code == 403)
    r = cv.get(reverse('evidencia_list'), {'page_size': 30})
    ok('Verificador: ve evidencias de ambas delegaciones', r.context['page_obj'].paginator.count ==
       Evidence.objects.count())

    # ---------------- 5. Los 4 CRUD (crear, editar, eliminar lógico, exportar) ----------------
    c = cliente('admin_centro')
    ana = Employee.objects.get(user__username='funcionario_centro')
    fecha = min(abierto.end_date, __import__('django.utils.timezone', fromlist=['x']).localdate())
    datos = {'number': 'VERIF-001', 'employee': ana.pk, 'period': abierto.pk, 'activity_type': centro_act.activity_type_id,
             'date': fecha.isoformat(), 'description': 'Actividad de verificación', 'evidence_code': 'EV-VERIF-001'}
    r = c.post(reverse('actividad_create'), datos)
    nueva_act = Activity.objects.filter(number='VERIF-001').first()
    ok('CRUD Actividades: crear', r.status_code == 302 and nueva_act is not None)
    r = c.post(reverse('actividad_create'), {**datos, 'number': 'VERIF-002', 'evidence_code': 'EV-VERIF-002',
                                             'date': (fecha.replace(year=fecha.year + 1)).isoformat()})
    ok('CRUD Actividades: fecha futura rechazada', r.status_code == 200 and 'date' in r.context['form'].errors)
    r = c.post(reverse('actividad_update', args=[nueva_act.pk]), {**datos, 'description': 'Editada'})
    nueva_act.refresh_from_db()
    ok('CRUD Actividades: editar', r.status_code == 302 and nueva_act.description == 'Editada')

    social = Activity.objects.filter(delegation__name='Delegación Centro', period=abierto,
                                     activity_type__category='social').exclude(social_cases__step_number=3).first()
    paso = next(p for p in (1, 2, 3) if not SocialCase.all_objects.filter(activity=social, step_number=p).exists())
    r = c.post(reverse('atencion_create'), {'activity': social.pk, 'step_number': paso, 'description': 'Gestión verif',
                                            'date': fecha.isoformat(), 'result': 'En seguimiento'})  # RN-012
    ok('CRUD Atenciones: crear', r.status_code == 302)
    no_social = Activity.objects.filter(delegation__name='Delegación Centro', period=abierto).exclude(
        activity_type__category='social').first()
    r = c.post(reverse('atencion_create'), {'activity': no_social.pk, 'step_number': 1, 'description': 'x'})
    ok('CRUD Atenciones: solo actividades de atención social', r.status_code == 200)

    r = c.post(reverse('evidencia_create'), {'activity': nueva_act.pk, 'description': 'Foto',
                                             'file': SimpleUploadedFile('foto.png', png(), 'image/png')})
    evi = Evidence.objects.filter(activity=nueva_act).first()
    ok('CRUD Evidencias: subir imagen', r.status_code == 302 and evi is not None and evi.file.name.endswith('.png'))
    ok('Archivos: nombre seguro (no se usa el del usuario)', evi and 'foto' not in evi.file.name)
    r = c.post(reverse('evidencia_create'), {'activity': nueva_act.pk, 'description': 'Falso',
                                             'file': SimpleUploadedFile('acta.pdf', b'no soy pdf')})
    ok('Archivos: contenido falso rechazado', r.status_code == 200 and 'file' in r.context['form'].errors)
    grande = SimpleUploadedFile('grande.pdf', b'%PDF-' + b'0' * (settings.EVIDENCIA_TAMANO_MAXIMO_MB * 1024 * 1024 + 1))
    r = c.post(reverse('evidencia_create'), {'activity': nueva_act.pk, 'description': 'Grande', 'file': grande})
    ok('Archivos: supera 2 MB, rechazado', r.status_code == 200 and 'file' in r.context['form'].errors)
    r = c.post(reverse('evidencia_create'), {'activity': nueva_act.pk, 'description': 'Exe',
                                             'file': SimpleUploadedFile('virus.exe', b'MZ')})
    ok('Archivos: extensión no permitida rechazada', r.status_code == 200)
    ok('Archivos: la imagen subida se puede abrir', Image.open(evi.file.path).verify() is None)

    r = c.post(reverse('compromiso_create'), {'title': 'Compromiso de verificación', 'responsible': ana.pk,
                                              'due_date': fecha.isoformat(), 'status': 'registered'})
    comp = Commitment.objects.filter(title='Compromiso de verificación').first()
    ok('CRUD Compromisos: crear (delegación la pone el sistema)', r.status_code == 302 and comp and comp.delegation == ana.delegation)
    r = c.post(reverse('compromiso_update', args=[comp.pk]), {'title': comp.title, 'responsible': ana.pk,
               'due_date': fecha.isoformat(), 'status': 'done', 'notes': ''})
    ok('CRUD Compromisos: realizado exige observaciones', r.status_code == 200)

    for nombre, url in [('Actividades', 'actividad_export'), ('Atenciones', 'atencion_export'),
                        ('Evidencias', 'evidencia_export'), ('Compromisos', 'compromiso_export')]:
        try:
            hoja = xlsx(c.get(reverse(url)))
            ok(f'Excel {nombre}: se abre y trae encabezados', hoja.max_row >= 1 and hoja['A1'].font.bold)
        except Exception as e:  # noqa: BLE001
            ok(f'Excel {nombre}: se abre', False, str(e))

    r = c.post(reverse('evidencia_delete', args=[evi.pk]))
    evi_bd = Evidence.all_objects.get(pk=evi.pk)
    ok('Borrado lógico: la evidencia queda con deleted_at y conserva su archivo',
       r.status_code == 302 and evi_bd.deleted_at and os.path.exists(evi_bd.file.path))
    Evidence.objects.create(activity=nueva_act, description='Hija', file='evidencias/x.png')
    c.post(reverse('actividad_delete', args=[nueva_act.pk]))
    nueva_act = Activity.all_objects.get(pk=nueva_act.pk)
    ok('Borrado lógico: la actividad y sus evidencias caen juntas',
       nueva_act.deleted_at and not Evidence.objects.filter(activity=nueva_act).exists())
    ok('Borrado lógico: no aparece en el listado', not c.get(reverse('actividad_list'), {'page_size': 30}).context[
        'page_obj'].paginator.object_list.filter(pk=nueva_act.pk).exists())
    ok('Borrado lógico: admin_centro ya no ve la actividad eliminada en el Admin',
       c.get(f'/admin/actividades/activity/{nueva_act.pk}/change/').status_code == 302 if nueva_act.deleted_at else True)
    sa = cliente('admin_sgr')
    ok('Regla de negocio: el superadmin ve lo que eliminó un admin de delegación (solo lectura)',
       nueva_act.deleted_at and sa.get(f'/admin/actividades/activity/{nueva_act.pk}/change/').status_code == 200
       and nueva_act in list(sa.get('/admin/actividades/activity/', {'registro': 'eliminados'}).context['cl'].result_list))
    ok('Regla de negocio: queda registrado quién eliminó',
       AuditLog.objects.filter(action='eliminar', entity_type='Activity', entity_id=nueva_act.pk, user__username='admin_centro').exists())
    nueva_act.restore()
    ok('Borrado lógico: restaurar recupera la actividad y las evidencias que cayeron con ella',
       Evidence.objects.filter(activity=nueva_act, description='Hija').exists())
    ok('Eliminar: solo por POST', c.get(reverse('actividad_delete', args=[nueva_act.pk])).status_code == 405)
    lista = c.get(reverse('actividad_list')).content.decode()
    ok('SweetAlert2: el listado carga la confirmación de borrado', 'vendor/sweetalert2/sweetalert2.all.min.js' in lista and 'confirmar.js' in lista)

    # ---------------- 6. Revisión de evidencias y estado de la actividad ----------------
    Activity.objects.filter(pk=nueva_act.pk).update(validation_status='pending')
    Evidence.objects.filter(activity=nueva_act).delete()
    e1 = Evidence.objects.create(activity=nueva_act, description='Borrosa', file='evidencias/x.png')
    cv = cliente('verificador_leia')
    cv.post(reverse('evidencia_update', args=[e1.pk]), {'status': 'rejected', 'result': 'No se lee'})
    nueva_act.refresh_from_db()
    ok('Revisión: rechazar deja la actividad rechazada', nueva_act.validation_status == 'rejected')
    ok('Revisión: queda Validation y traza de auditoría', Validation.objects.filter(evidence=e1, status='rejected').exists()
       and AuditLog.objects.filter(entity_id=e1.pk, action='evidencia_rejected').exists())
    e2 = Evidence.objects.create(activity=nueva_act, description='Corregida', file='evidencias/y.png')
    nueva_act.refresh_from_db()
    ok('Revisión: subir una corrección deja la actividad pendiente', nueva_act.validation_status == 'pending')
    cv.post(reverse('evidencia_update', args=[e2.pk]), {'status': 'approved', 'result': 'Correcta'})
    nueva_act.refresh_from_db()
    ok('Revisión: aprobar la corrección aprueba la actividad', nueva_act.validation_status == 'approved')
    cf = cliente('funcionario_centro')
    ok('Revisión: aprobada, el funcionario ya no la edita', cf.get(reverse('actividad_update', args=[nueva_act.pk])).status_code == 403)

    # ---------------- 7. Período cerrado ----------------
    # De Centro: con una de Norte, admin_centro recibe 404 (fuera de su delegación) y no se probaría el cierre.
    historica = Activity.all_objects.filter(period=cerrado, delegation__name='Delegación Centro').first()
    if historica:
        ok('Período cerrado: ni el admin edita (403)', c.get(reverse('actividad_update', args=[historica.pk])).status_code == 403)
    r = c.post(reverse('actividad_create'), {**datos, 'number': 'VERIF-003', 'evidence_code': 'EV-VERIF-003', 'period': cerrado.pk})
    ok('Período cerrado: no se registran actividades nuevas', r.status_code == 200)

    # Gestión social de una actividad aprobada: el funcionario ya no la edita.
    gestion = SocialCase.objects.filter(activity__employee__user__username='funcionario_centro',
                                        activity__period=abierto).first()
    if gestion:
        Activity.objects.filter(pk=gestion.activity_id).update(validation_status='approved')
        ok('Actividad aprobada: el funcionario tampoco edita sus gestiones sociales (403)',
           cliente('funcionario_centro').get(reverse('atencion_update', args=[gestion.pk])).status_code == 403)

    # ---------------- 8. Dashboard según el rol ----------------
    for usuario in ('admin_sgr', 'admin_centro', 'funcionario_centro', 'verificador_leia'):
        cu = cliente(usuario)
        for k in cu.get(reverse('dashboard')).context['indicadores']:
            if k.get('url'):
                filas = cu.get(k['url']).context['page_obj'].paginator.count
                ok(f'Dashboard {usuario}: «{k["etiqueta"]}» ({k["valor"]}) coincide con su lista', filas == k['valor'],
                   f'lista {filas}')

    for usuario, debe, no_debe in [
        ('admin_sgr', ['Resumen de gestión', 'Delegación Centro', 'Delegación Norte', 'a hoy'], []),
        ('admin_centro', ['Resumen de gestión', 'Resumen por funcionario · Delegación Centro'], ['Delegación Norte']),
        ('funcionario_centro', ['Mi gestión', 'Actividades rechazadas', 'Mi avance por tipo'], ['Resumen por funcionario', 'de su equipo']),
        ('verificador_leia', ['Revisión de evidencias', 'Estado de evidencias', 'status=pending'], ['Compromisos vencidos', 'Resumen por funcionario']),
    ]:
        html = cliente(usuario).get(reverse('dashboard')).content.decode()
        faltan = [t for t in debe if t not in html] + [f'NO:{t}' for t in no_debe if t in html]
        ok(f'Dashboard {usuario}: contenido según su rol', not faltan, ', '.join(faltan))
    r = cliente('admin_centro').get(reverse('dashboard'), {'period': cerrado.pk})
    ok('Dashboard: período sin metas lo informa', 'Sin metas definidas' in r.content.decode())
    html = cliente('admin_centro').get(reverse('dashboard')).content.decode()
    ok('Dashboard: sin coma decimal en CSS/SVG', not re.search(r'(width|left|height): \d+,\d|dash(array|offset)="[^"]*\d,\d', html))

    # ---------------- 9. Admin integrado ----------------
    html = cliente('admin_centro').get('/admin/').content.decode()
    ok('Admin: misma barra y menú del sistema', 'SGR · La Serena' in html and 'sgr-menu-boton' in html)
    html = cliente('admin_centro').get('/admin/actividades/activity/').content.decode()
    ok('Admin: en español', 'Seleccione una opción' in html and 'Select an option' not in html)

    # ---------------- 10. Volumen de datos, nombres y seguridad del repositorio ----------------
    negocio = (Activity.objects.count() + SocialCase.objects.count() + Evidence.objects.count() + Validation.objects.count()
               + Commitment.objects.count() + CommitmentFollowUp.objects.count() + Employee.objects.count())
    ok('Volumen: 1.000 o más registros de negocio', negocio >= 1000, str(negocio))
    for d in ('Delegación Centro', 'Delegación Norte'):
        ok(f'Volumen: {d} con actividades y evidencias',
           Activity.objects.filter(delegation__name=d).count() > 200 and Evidence.objects.filter(activity__delegation__name=d).count() > 100)
    tablas = connection.introspection.table_names()
    ok('Nombres en inglés: tablas del negocio', all(t in tablas for t in ['activity', 'evidence', 'validation', 'commitment', 'employee', 'goal']))
    archivos = subprocess.run(['git', 'ls-files'], capture_output=True, text=True).stdout.split()
    ok('Secretos: .env fuera del repositorio', '.env' not in archivos and '.env.example' in archivos)
    sospechosos = subprocess.run(['git', 'grep', '-nE', r"SECRET_KEY\s*=\s*['\"][^'\"]{20,}|Admin#2026SGR|AdminCentro#2026", '--', '.', ':!scripts/verificacion_e2e.py'],
                                 capture_output=True, text=True).stdout.strip()
    ok('Secretos: sin SECRET_KEY ni contraseñas demo antiguas en el código', not sospechosos, sospechosos[:200])
    ok('Configuración: DEBUG apagado por defecto', "os.getenv('DEBUG', 'False')" in open('config/settings.py').read())
    ok('Despliegue: archivos de AWS presentes', all(os.path.exists(f'deploy/{f}') for f in ('setup_ec2.sh', 'gunicorn-sgr.service', 'nginx-sgr.conf')))

    transaction.set_rollback(True)

shutil.rmtree(media, ignore_errors=True)  # carpeta temporal de archivos subidos durante la verificación

fallas = [r for r in RESULTADOS if not r[1]]
for punto, bien, detalle in RESULTADOS:
    print(('OK   ' if bien else 'FALLA') + ' ' + punto + (f'  [{detalle}]' if detalle and not bien else ''))
print(f'\n{len(RESULTADOS) - len(fallas)}/{len(RESULTADOS)} verificaciones correctas')
